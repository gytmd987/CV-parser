"""워드 파일 읽기 · 쓰기.

## 읽기

검증 자료 한 사람 = 제목(이름 직책) · (소속) · 2열 표(항목 | 검증 의견) · 작성자.
예시 결과 파일은 이것이 여러 사람 이어 붙은 모양일 수 있어서, 표를 기준으로
앞뒤 문단을 갈라 사람마다 묶는다(`항목들`).

## 쓰기

결과 파일은 **예시 결과 파일을 틀로 쓴다.** 글꼴 · 크기 · 칸 너비 · 테두리 ·
여백을 코드로 다시 만들면 반드시 어긋난다. 예시의 문단 · 표 XML 을 그대로 복사해
글자만 바꿔 넣는다 (`한장_만들기` · `요약_만들기`).
"""

from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

W_P, W_TBL, W_SECT = qn("w:p"), qn("w:tbl"), qn("w:sectPr")


class 형식오류(ValueError):
    """파일이 기대한 모양이 아니다. 화면에 그대로 보여 줄 말."""


@dataclass
class 검증자료:
    제목: str                         # "홍길동 책임연구원" — 원문 그대로
    소속: str                         # 괄호 안
    항목: list[tuple[str, str]]       # [(항목명, 검증 의견)]
    끝글: str = ""                    # 표 뒤 문단 (작성자가 들어 있다)
    파일명: str = ""

    @property
    def 이름(self) -> str:
        return 이름_직책(self.제목)[0]

    @property
    def 직책(self) -> str:
        return 이름_직책(self.제목)[1]


def 이름_직책(제목: str) -> tuple[str, str]:
    조각 = [x for x in re.split(r"[\s,·/]+", 제목.strip()) if x]
    if not 조각:
        return "", ""
    return 조각[0], " ".join(조각[1:])


# --- 읽기 -------------------------------------------------------------------------
def _열기(내용: bytes | str):
    try:
        return Document(io.BytesIO(내용) if isinstance(내용, bytes) else 내용)
    except Exception as exc:  # noqa: BLE001 - 무엇이든 사람이 읽을 말로
        raise 형식오류("워드(.docx) 파일을 열지 못했습니다. .doc 이면 .docx 로 "
                    f"다시 저장해 주세요. ({type(exc).__name__})") from exc


def _문단글(el, 부모) -> str:
    return Paragraph(el, 부모).text.strip()


def 칸글(tc_or_cell) -> str:
    """칸 안의 문단을 줄바꿈으로 잇는다. 빈 줄은 뺀다."""
    tc = getattr(tc_or_cell, "_tc", tc_or_cell)
    줄 = ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in tc.iter(W_P)]
    return "\n".join(x.strip() for x in 줄 if x.strip())


def _표_줄(tbl_el, 부모) -> list[list[str]]:
    """표를 줄 × 칸 글로. 병합된 칸이 여러 번 나오는 것은 한 번만."""
    표 = Table(tbl_el, 부모)
    줄들 = []
    for row in 표.rows:
        본것, 칸들 = set(), []
        for c in row.cells:
            if id(c._tc) in 본것:
                continue
            본것.add(id(c._tc))
            칸들.append(칸글(c))
        줄들.append(칸들)
    return 줄들


def _머리줄인가(칸들: list[str]) -> bool:
    합 = "".join(칸들).replace(" ", "")
    return "항목" in 합 and ("의견" in 합 or "검증" in 합)


@dataclass
class _조각:
    앞문단: list = field(default_factory=list)    # (el, 글)
    표들: list = field(default_factory=list)       # tbl el
    뒤문단: list = field(default_factory=list)


def 항목들(문서) -> list[검증자료]:
    """문서 안의 사람들. 표(이어진 2열 표 묶음) 하나가 한 사람이다."""
    body = 문서.element.body
    부모 = 문서._body
    조각들: list[_조각] = []
    지금 = _조각()
    for el in body.iterchildren():
        if el.tag == W_P:
            글 = _문단글(el, 부모)
            if not 글:
                continue
            (지금.뒤문단 if 지금.표들 else 지금.앞문단).append((el, 글))
        elif el.tag == W_TBL:
            줄 = _표_줄(el, 부모)
            열수 = max((len(r) for r in 줄), default=0)
            if 열수 < 2:
                continue
            이어짐 = 지금.표들 and not 지금.뒤문단 and not (줄 and _머리줄인가(줄[0]))
            if 지금.표들 and not 이어짐:
                새 = _조각(앞문단=지금.뒤문단)
                지금.뒤문단 = []           # 다음 사람 제목까지 끝글로 먹지 않게
                조각들.append(지금)
                지금 = 새
            지금.표들.append(el)
    if 지금.표들:
        조각들.append(지금)
    elif 조각들:
        조각들[-1].뒤문단 += 지금.앞문단

    사람들 = []
    for i, 조 in enumerate(조각들):
        앞 = [g for _e, g in 조.앞문단]
        제목, 소속, 쓴수 = _제목_소속(앞)
        # 앞문단 중 제목 앞의 것은 앞 사람의 끝글이다
        if i > 0 and 쓴수 < len(앞):
            남은 = 앞[: len(앞) - 쓴수]
            사람들[-1].끝글 = "\n".join(filter(None, [사람들[-1].끝글, *남은]))
        항목 = []
        for el in 조.표들:
            for 칸들 in _표_줄(el, 부모):
                if _머리줄인가(칸들) or not any(칸들):
                    continue
                if len(칸들) == 1:
                    continue
                항목.append((칸들[0], "\n".join(c for c in 칸들[1:] if c)))
        사람들.append(검증자료(제목=제목, 소속=소속, 항목=항목,
                         끝글="\n".join(g for _e, g in 조.뒤문단)))
    return 사람들


def _제목_소속(앞: list[str]) -> tuple[str, str, int]:
    """표 앞 문단에서 제목과 (소속). 소속은 괄호로 싸인 줄. (제목, 소속, 쓴 문단 수)"""
    for i in range(len(앞) - 1, -1, -1):
        줄 = 앞[i].strip()
        m = re.fullmatch(r"[(（](.*)[)）]", 줄)
        if m and i > 0:
            return 앞[i - 1].strip(), m.group(1).strip(), len(앞) - (i - 1)
        # 제목과 소속이 한 문단에 줄바꿈으로 들어 있는 경우
        if "\n" in 줄:
            a, _, b = 줄.partition("\n")
            m = re.fullmatch(r"[(（](.*)[)）]", b.strip())
            if m:
                return a.strip(), m.group(1).strip(), len(앞) - i
    if 앞:
        return 앞[-1].strip(), "", 1
    return "", "", 0


def 검증자료_읽기(내용: bytes | str, 파일명: str = "") -> 검증자료:
    사람들 = 항목들(_열기(내용))
    if not 사람들:
        raise 형식오류("2열 표(항목 | 검증 의견)를 찾지 못했습니다.")
    사람 = 사람들[0]
    if not 사람.제목:
        raise 형식오류("표 위에서 제목(이름 직책)을 찾지 못했습니다.")
    if len(사람.항목) < 2:
        raise 형식오류(f"표에서 항목을 {len(사람.항목)}개밖에 못 읽었습니다.")
    사람.파일명 = 파일명
    return 사람


def 요약표_읽기(내용: bytes | str) -> tuple[str, list[str], list[list[str]]]:
    """(제목, 머리줄, 데이터 줄들)."""
    문서 = _열기(내용)
    제목 = next((p.text.strip() for p in 문서.paragraphs if p.text.strip()), "")
    if not 문서.tables:
        raise 형식오류("요약 파일에서 표를 찾지 못했습니다.")
    줄 = _표_줄(문서.tables[0]._tbl, 문서._body)
    return 제목, 줄[0], [r for r in 줄[1:] if any(r)]


# --- 작성자 ------------------------------------------------------------------------
_안이름 = set("""작성자 작성 검증자 검증 검증위원 위원 위원장 서명 날인 일자 날짜 일시 소속 성명
직위 직책 수석 책임 선임 전임 원급 연구원 팀장 부장 차장 과장 대리 사원 교수 박사 확인
담당 담당자 및 이상 이하 귀하 년 월 일""".split())


def 작성자_찾기(끝글: str) -> list[str]:
    """끝글에서 사람 이름(한글 2~4자)만. 직함 · '작성자' 같은 낱말은 뺀다."""
    이름 = []
    for 낱말 in re.findall(r"[가-힣]{2,4}", 끝글):
        if 낱말 in _안이름 or 낱말.endswith(("연구원", "위원")):
            continue
        if 낱말 not in 이름:
            이름.append(낱말)
    return 이름


# --- 칸 너비 · 글자 크기 ---------------------------------------------------------------
@dataclass
class 칸모양:
    폭_pt: float      # 글이 들어가는 너비 (칸 너비 - 좌우 여백)
    크기_pt: float    # 글자 크기
    글꼴: str = ""

    @property
    def 한줄_em(self) -> float:
        return self.폭_pt / self.크기_pt


def _twip(값) -> float:
    try:
        return float(값) / 20.0
    except (TypeError, ValueError):
        return 0.0


def _기본크기(문서) -> float:
    for rpr in 문서.styles.element.iter(qn("w:rPrDefault")):
        sz = rpr.find(".//" + qn("w:sz"))
        if sz is not None:
            return float(sz.get(qn("w:val"))) / 2
    try:
        크기 = 문서.styles["Normal"].font.size
        if 크기:
            return 크기.pt
    except KeyError:
        pass
    return 10.0


def _칸글꼴(tc, 문서) -> tuple[float, str]:
    for r in tc.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            continue
        sz = rpr.find(qn("w:sz"))
        fonts = rpr.find(qn("w:rFonts"))
        글꼴 = ""
        if fonts is not None:
            글꼴 = fonts.get(qn("w:eastAsia")) or fonts.get(qn("w:ascii")) or ""
        if sz is not None:
            return float(sz.get(qn("w:val"))) / 2, 글꼴
    return _기본크기(문서), ""


def 칸모양_읽기(tbl_el, 열: int, 문서, 기본폭_cm: float = 0.0, 기본크기: float = 0.0) -> 칸모양:
    """표의 한 열에 글이 들어갈 너비와 글자 크기."""
    폭 = 0.0
    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        if 열 < len(cols):
            폭 = _twip(cols[열].get(qn("w:w")))
    줄 = tbl_el.findall(qn("w:tr"))
    tc = None
    if len(줄) > 1:
        tcs = 줄[1].findall(qn("w:tc"))
        if 열 < len(tcs):
            tc = tcs[열]
            w = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
            if w is not None and w.get(qn("w:type")) in (None, "dxa"):
                폭 = _twip(w.get(qn("w:w"))) or 폭
    if not 폭 and 기본폭_cm:
        폭 = 기본폭_cm / 2.54 * 72
    여백 = 0.0
    mar = tbl_el.find(qn("w:tblPr") + "/" + qn("w:tblCellMar"))
    for 쪽 in ("left", "right", "start", "end"):
        el = mar.find(qn(f"w:{쪽}")) if mar is not None else None
        if el is not None:
            여백 += _twip(el.get(qn("w:w")))
    if mar is None or not 여백:
        여백 = 2 * 5.4                    # 워드 기본 0.19cm
    크기, 글꼴 = _칸글꼴(tc, 문서) if tc is not None else (0.0, "")
    return 칸모양(폭_pt=max(폭 - 여백, 20.0), 크기_pt=크기 or 기본크기 or 10.0, 글꼴=글꼴)


# --- 쓰기 -------------------------------------------------------------------------
def 글넣기(p_el, 글: str) -> None:
    """문단의 글을 바꾼다. 첫 글자 모양(rPr)을 그대로 쓴다."""
    첫 = p_el.find(qn("w:r"))
    rpr = copy.deepcopy(첫.find(qn("w:rPr"))) if 첫 is not None and 첫.find(qn("w:rPr")) is not None else None
    for 자식 in list(p_el):
        if 자식.tag != qn("w:pPr"):
            p_el.remove(자식)
    r = p_el.makeelement(qn("w:r"), {})
    if rpr is not None:
        r.append(rpr)
    for i, 줄 in enumerate(글.split("\n")):
        if i:
            r.append(r.makeelement(qn("w:br"), {}))
        t = r.makeelement(qn("w:t"), {})
        t.text = 줄
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
    p_el.append(r)


def 칸넣기(tc, 글: str) -> None:
    """칸 글을 바꾼다. 줄마다 문단 하나 (첫 문단 모양을 복사)."""
    문단들 = tc.findall(W_P)
    틀 = copy.deepcopy(문단들[0]) if 문단들 else tc.makeelement(W_P, {})
    for p in 문단들:
        tc.remove(p)
    for 줄 in (글.split("\n") if 글 else [""]):
        새 = copy.deepcopy(틀)
        글넣기(새, 줄)
        tc.append(새)


def _쪽나눔(p_el) -> None:
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = p_el.makeelement(qn("w:pPr"), {})
        p_el.insert(0, ppr)
    if ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.insert(0, ppr.makeelement(qn("w:pageBreakBefore"), {}))


def _몸통비우기(문서) -> None:
    body = 문서.element.body
    for el in list(body):
        if el.tag != W_SECT:
            body.remove(el)


def _몸통에(문서, el) -> None:
    body = 문서.element.body
    sect = body.find(W_SECT)
    if sect is not None:
        sect.addprevious(el)
    else:
        body.append(el)


@dataclass
class 한장결과:
    제목: str
    소속: str
    항목: list[tuple[str, str]]
    작성자: list[str]


def _작성자글(틀글: str, 이름들: list[str]) -> str:
    """예시의 작성자 줄 모양대로 이름만 바꾼다 ('작성자 김OO / 이OO / 박OO')."""
    옛 = 작성자_찾기(틀글)
    if 옛 and len(옛) == len(이름들):
        글 = 틀글
        for a, b in zip(옛, 이름들):
            글 = 글.replace(a, b, 1)
        return 글
    머리 = "작성자 "
    if "작성자" in 틀글:
        머리 = 틀글[: 틀글.index("작성자") + 3] + " "
    return 머리 + " / ".join(이름들)


def 한장_만들기(틀: bytes, 결과들: list[한장결과]) -> bytes:
    """예시 1장 파일의 첫 사람 모양을 틀로, 여러 사람을 한 파일에 (사람마다 새 쪽)."""
    문서 = _열기(틀)
    body = 문서.element.body
    자식 = [el for el in body if el.tag != W_SECT]
    표위치 = next((i for i, el in enumerate(자식) if el.tag == W_TBL), None)
    if 표위치 is None:
        raise 형식오류("1장 예시 파일에서 표를 찾지 못했습니다.")
    부모 = 문서._body
    # 제목 · 소속: 표 앞의 글 있는 문단 (가장 가까운 두 개)
    앞 = [i for i in range(표위치) if 자식[i].tag == W_P and _문단글(자식[i], 부모)]
    제목i = 앞[-2] if len(앞) >= 2 else (앞[-1] if 앞 else None)
    소속i = 앞[-1] if len(앞) >= 2 else None
    # 작성자: 표 뒤 첫 글 있는 문단
    뒤 = [i for i in range(표위치 + 1, len(자식))
         if 자식[i].tag == W_P and _문단글(자식[i], 부모)]
    작성자i = 뒤[0] if 뒤 else None
    시작 = 제목i if 제목i is not None else 표위치
    끝 = 작성자i if 작성자i is not None else 표위치
    틀조각 = [copy.deepcopy(el) for el in 자식[시작: 끝 + 1]]
    작성자틀글 = _문단글(자식[작성자i], 부모) if 작성자i is not None else "작성자"

    _몸통비우기(문서)
    for n, 결과 in enumerate(결과들):
        조각 = [copy.deepcopy(el) for el in 틀조각]
        for k, el in enumerate(조각):
            원래 = 시작 + k
            if 원래 == 제목i:
                글넣기(el, 결과.제목)
                if n:
                    _쪽나눔(el)
            elif 원래 == 소속i:
                글넣기(el, f"({결과.소속})" if 결과.소속 else "")
            elif 원래 == 작성자i:
                글넣기(el, _작성자글(작성자틀글, 결과.작성자))
            elif 원래 == 표위치:
                _표채우기(el, 결과.항목)
        if n and 제목i is None:
            _쪽나눔(조각[0]) if 조각[0].tag == W_P else None
        for el in 조각:
            _몸통에(문서, el)
    나감 = io.BytesIO()
    문서.save(나감)
    return 나감.getvalue()


def _표채우기(tbl_el, 항목: list[tuple[str, str]]) -> None:
    줄 = tbl_el.findall(qn("w:tr"))
    머리 = 1 if 줄 and _머리줄인가([칸글(tc) for tc in 줄[0].findall(qn("w:tc"))]) else 0
    본줄 = 줄[머리:]
    if not 본줄:
        raise 형식오류("1장 예시 표에 내용 줄이 없습니다.")
    for tr in 본줄:
        tbl_el.remove(tr)
    for i, (이름, 의견) in enumerate(항목):
        tr = copy.deepcopy(본줄[min(i, len(본줄) - 1)])
        tcs = tr.findall(qn("w:tc"))
        칸넣기(tcs[0], 이름)
        if len(tcs) > 1:
            칸넣기(tcs[1], 의견)
        tbl_el.append(tr)


def 요약_만들기(틀: bytes, 줄들: list[list[str]]) -> bytes:
    """예시 요약 파일의 표 첫 내용 줄을 틀로 사람 수만큼 줄을 만든다."""
    문서 = _열기(틀)
    if not 문서.tables:
        raise 형식오류("요약 예시 파일에서 표를 찾지 못했습니다.")
    tbl = 문서.tables[0]._tbl
    trs = tbl.findall(qn("w:tr"))
    if len(trs) < 2:
        raise 형식오류("요약 예시 표에 내용 줄이 없습니다.")
    틀줄 = trs[1]
    for tr in trs[1:]:
        tbl.remove(tr)
    for 값들 in 줄들:
        tr = copy.deepcopy(틀줄)
        for tc, 값 in zip(tr.findall(qn("w:tc")), 값들):
            칸넣기(tc, 값)
        tbl.append(tr)
    나감 = io.BytesIO()
    문서.save(나감)
    return 나감.getvalue()
