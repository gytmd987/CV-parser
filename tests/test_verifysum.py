"""검증 자료 정리 (verifysum) — 워드 읽기 · 줄 수 세기 · 결과 파일 · 웹 흐름."""

from __future__ import annotations

import io
import json
import re
import threading
import time
import urllib.parse
import urllib.request
import uuid
from http.server import ThreadingHTTPServer

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from verifysum import docxio as D
from verifysum import fit, generate, learn

항목명 = ["연구 전문성", "주요 성과", "연구 수행 역량", "협업·소통", "연구 윤리", "성장 가능성", "종합 의견"]


def _바탕(run, 크기=11):
    run.font.size = Pt(크기)
    run.font.name = "바탕체"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "바탕체")


def 검증파일(이름="홍길동", 직책="책임연구원", 소속="한국과학기술원 신소재공학과",
          의견=None, 작성자="작성자: 김철수 수석 / 이영희 책임 / 박민수 선임", 사람수=1) -> bytes:
    doc = Document()
    for n in range(사람수):
        이 = 이름 if 사람수 == 1 else f"{이름}{n}"
        p = doc.add_paragraph()
        if n:
            p.add_run().add_break(WD_BREAK.PAGE)
        r = p.add_run(f"{이} {직책}")
        r.bold = True
        doc.add_paragraph(f"({소속})")
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "검증 의견"
        for i, 항 in enumerate(항목명):
            c = t.add_row().cells
            c[0].text = 항
            c[1].text = (의견 or "원자층 식각 공정을 오랫동안 연구하였으며 관련 논문을 다수 발표하였다. " * 4)
            c[1].add_paragraph(f"추가 의견 {i}: 동료 평가가 좋고 성실함.")
        doc.add_paragraph("")
        doc.add_paragraph(작성자 if 사람수 == 1 else "작성자 " + " / ".join(["김민수 이지은 박서준".split(), "최유진 정하늘 강도윤".split(), "윤서아 장민호 임채원".split()][n]))
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def 한장파일() -> bytes:
    doc = Document()
    doc.add_paragraph().add_run("홍길동 책임연구원").bold = True
    doc.add_paragraph("(한국과학기술원 신소재공학과)")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "검증 의견"
    for 항 in 항목명:
        c = t.add_row().cells
        c[0].text = 항
        c[1].text = "원자층 식각 공정 연구 경력 10년, 관련 논문 다수 발표함."
    doc.add_paragraph("작성자 김철수 / 이영희 / 박민수")
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def 요약파일() -> bytes:
    doc = Document()
    doc.add_paragraph("개인별 검증결과 요약")
    t = doc.add_table(rows=2, cols=4)
    t.autofit = False
    for i, (머리, 폭) in enumerate(zip(["후보자", "전문성/성과", "연구전문가 역량", "결과"],
                                    [Cm(2.5), Cm(7), Cm(7), Cm(2)])):
        t.columns[i].width = 폭
        for row in t.rows:
            row.cells[i].width = 폭
        t.rows[0].cells[i].text = 머리
    값 = ["홍길동", "원자층 식각·플라즈마 공정 전문, 고선택비 ALE 공정 개발", "성실하고 협업에 적극적인 연구 태도", "추천"]
    for i, v in enumerate(값):
        cell = t.rows[1].cells[i]
        _바탕(cell.paragraphs[0].add_run(v))
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


# --- 읽기 -------------------------------------------------------------------------
def test_검증_자료를_읽는다():
    자료 = D.검증자료_읽기(검증파일(), "a.docx")
    assert 자료.제목 == "홍길동 책임연구원" and 자료.이름 == "홍길동" and 자료.직책 == "책임연구원"
    assert 자료.소속 == "한국과학기술원 신소재공학과"
    assert [a for a, _b in 자료.항목] == 항목명
    assert "추가 의견 0" in 자료.항목[0][1] and "\n" in 자료.항목[0][1]
    assert D.작성자_찾기(자료.끝글) == ["김철수", "이영희", "박민수"]


def test_여러_사람이_든_파일을_사람마다_가른다():
    사람들 = D.항목들(Document(io.BytesIO(검증파일(사람수=3))))
    assert [p.제목 for p in 사람들] == ["홍길동0 책임연구원", "홍길동1 책임연구원", "홍길동2 책임연구원"]
    assert all(len(p.항목) == 7 for p in 사람들)
    assert [D.작성자_찾기(p.끝글) for p in 사람들] == [
        ["김민수", "이지은", "박서준"], ["최유진", "정하늘", "강도윤"], ["윤서아", "장민호", "임채원"]]
    assert "홍길동1" not in 사람들[0].끝글               # 다음 사람 제목을 끝글로 먹지 않는다


def test_docx_가_아니면_알기_쉽게():
    with pytest.raises(D.형식오류, match="docx"):
        D.검증자료_읽기(b"not a word file")


def test_요약칸_너비와_글자크기():
    문서 = Document(io.BytesIO(요약파일()))
    칸 = D.칸모양_읽기(문서.tables[0]._tbl, 1, 문서)
    assert 칸.크기_pt == 11 and 칸.글꼴 == "바탕체"
    assert 16.5 < 칸.한줄_em < 17.5                 # 7cm - 여백, 11pt → 한글 약 17자


# --- 줄 수 -------------------------------------------------------------------------
def test_줄수_세기():
    assert fit.줄수("가" * 17, 17.6) == 1
    assert fit.줄수("가" * 18, 17.6) == 2
    assert fit.줄수("abcd " * 6, 17.6) == 1          # 영문은 반 칸
    assert fit.줄수("한 줄\n두 줄", 17.6) == 2
    assert fit.줄수("가나다라마 " * 7, 17.6) == 3      # 띄어쓰기에서만 끊는다고 본다
    잘린 = fit.자르기("가나다라마 " * 9, 17.6)
    assert fit.줄수(잘린, 17.6) <= 2 and 잘린


# --- 학습 · 생성 ---------------------------------------------------------------------
class 가짜LLM:
    """처음엔 요약 칸을 길게 쓰고, 고치라고 하면 줄인다."""

    def __init__(self, 늘_길게=False):
        self.호출 = []
        self.늘_길게 = 늘_길게

    def chat_text(self, messages, **kw):
        self.호출.append(messages)
        return "[문체]\n- ~함 으로 끝맺는다."

    def chat_json(self, messages, schema, **kw):
        self.호출.append(messages)
        n = schema["properties"]["검증의견"]["minItems"]
        고침 = messages[-1]["role"] == "user" and "고쳐" in messages[-1]["content"]
        긴 = "원자층 식각과 플라즈마 공정 전반에 걸친 깊은 전문성과 다수의 국제 학술지 논문 실적을 갖춤"
        짧 = "원자층 식각 공정 전문, 고선택비 ALE 개발"
        return {"검증의견": [f"의견 {i}: 핵심만 정리함." for i in range(n)],
                "전문성_성과": 긴 if (self.늘_길게 or not 고침) else 짧,
                "연구전문가_역량": "성실하고 협업에 적극적임",
                "작성자": ["김철수", "이영희", "박민수"]}

    def close(self):
        pass


@pytest.fixture()
def 학습():
    h = learn.분석(검증파일(), 한장파일(), 요약파일())
    h["지침"] = learn.기본지침(h)
    return h


def test_예시를_분석한다(학습):
    assert 학습["원본"]["제목"] == "홍길동 책임연구원"
    assert len(학습["한장"]["항목"]) == 7 and 학습["분량"]["표높이"] > 0
    assert 학습["요약"]["열"] == {"후보자": 0, "전문성": 1, "역량": 2, "결과": 3}
    assert 학습["칸"]["전문성"]["크기_pt"] == 11
    assert learn.지침_만들기(학습, 가짜LLM()).startswith("[문체]")


def test_지침_실패해도_기본_지침(학습):
    class 죽은LLM:
        def chat_text(self, *a, **k):
            raise RuntimeError("연결 안 됨")
    assert "[문체]" in learn.지침_만들기(학습, 죽은LLM())


def test_요약_칸이_길면_고쳐_쓰게_한다(학습):
    llm = 가짜LLM()
    결 = generate.한사람(D.검증자료_읽기(검증파일(이름="김민준")), 학습, llm)
    assert len(llm.호출) == 2                                # 한 번 고쳐 씀
    assert 결.전문성 == "원자층 식각 공정 전문, 고선택비 ALE 개발" and not 결.경고
    지시 = llm.호출[0][0]["content"]
    assert "두 줄 안" in 지시 and "지어내지" in 지시 and "약 " in 지시


def test_끝까지_길면_두_줄에_맞춰_자른다(학습):
    결 = generate.한사람(D.검증자료_읽기(검증파일()), 학습, 가짜LLM(늘_길게=True))
    assert fit.줄수(결.전문성, 학습["칸"]["전문성"]["폭_pt"] / 11) <= 2
    assert any("잘랐습니다" in w for w in 결.경고)


def test_결과_파일_두_개(학습):
    llm = 가짜LLM()
    결과들 = [generate.한사람(D.검증자료_읽기(검증파일(이름=n)), 학습, llm) for n in ("가나다", "라마바")]
    한장 = Document(io.BytesIO(generate.한장파일(결과들, 한장파일())))
    assert len(한장.tables) == 2 and all(len(t.rows) == 8 for t in 한장.tables)
    글 = [p.text for p in 한장.paragraphs if p.text.strip()]
    assert 글[:3] == ["가나다 책임연구원", "(한국과학기술원 신소재공학과)", "작성자 김철수 / 이영희 / 박민수"]
    assert 한장.tables[1].rows[1].cells[1].text == "의견 0: 핵심만 정리함."
    둘째제목 = next(p for p in 한장.paragraphs if p.text == "라마바 책임연구원")
    assert 둘째제목._p.find(qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None
    assert 한장.paragraphs[0].runs[0].bold                   # 예시 모양(굵게)을 그대로

    요약 = Document(io.BytesIO(generate.요약파일(결과들, 요약파일(), 학습)))
    표 = 요약.tables[0]
    assert [c.text for c in 표.rows[0].cells] == ["후보자", "전문성/성과", "연구전문가 역량", "결과"]
    assert [r.cells[0].text for r in 표.rows[1:]] == ["가나다", "라마바"]
    assert all(r.cells[3].text == "추천" for r in 표.rows[1:])
    run = 표.rows[1].cells[1].paragraphs[0].runs[0]
    assert run.font.size == Pt(11) and run._element.rPr.rFonts.get(qn("w:eastAsia")) == "바탕체"
    assert 요약.paragraphs[0].text == "개인별 검증결과 요약"


# --- 웹 ----------------------------------------------------------------------------
def _멀티(필드: dict[str, list[tuple[str, bytes]]]) -> tuple[bytes, str]:
    경계 = uuid.uuid4().hex
    몸 = b""
    for 이름, 파일들 in 필드.items():
        for 파일명, 내용 in 파일들:
            몸 += (f"--{경계}\r\nContent-Disposition: form-data; name=\"{이름}\"; "
                  f"filename=\"{파일명}\"\r\nContent-Type: application/octet-stream\r\n\r\n").encode() + 내용 + b"\r\n"
    return 몸 + f"--{경계}--\r\n".encode(), f"multipart/form-data; boundary={경계}"


def test_웹으로_처음부터_끝까지(tmp_path, monkeypatch):
    from verifysum import web

    monkeypatch.setattr(web, "폴더", tmp_path)
    monkeypatch.setattr(web, "_클라이언트", lambda: 가짜LLM())
    srv = ThreadingHTTPServer(("127.0.0.1", 0), web.Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    U = f"http://127.0.0.1:{srv.server_address[1]}"

    class 안따라감(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k):
            return None
    opener = urllib.request.build_opener(안따라감)

    def post(경로, 몸, 형식):
        try:
            opener.open(urllib.request.Request(U + 경로, data=몸, headers={"Content-Type": 형식}))
        except urllib.error.HTTPError as e:
            return urllib.parse.unquote(e.headers.get("Location", ""))
        return ""

    try:
        assert "결과 만들기" in urllib.request.urlopen(U + "/").read().decode()
        위치 = post("/learn", *_멀티({"원본": [("a.docx", 검증파일())], "한장": [("b.docx", 한장파일())],
                                     "요약": [("c.docx", 요약파일())]}))
        assert "분석했습니다" in 위치
        첫 = urllib.request.urlopen(U + "/").read().decode()
        assert "홍길동 책임연구원" in 첫 and "[문체]" in 첫 and "한글 약 17자" in 첫

        post("/guide", urllib.parse.urlencode({"지침": "[문체]\n- 고친 지침"}).encode(),
             "application/x-www-form-urlencoded")
        assert learn.불러오기(tmp_path)["지침"] == "[문체]\n- 고친 지침"

        파일 = [(f"{i}.docx", 검증파일(이름=f"후보{i}")) for i in (10, 2, 1)] + [("깨짐.docx", b"x")]
        위치 = post("/run", *_멀티({"files": 파일}))
        assert "4개 파일" in 위치
        for _ in range(100):
            s = json.loads(urllib.request.urlopen(U + "/status").read())
            if not s["돌고있음"]:
                break
            time.sleep(0.05)
        assert [x["파일"] for x in s["줄"]] == ["1.docx", "2.docx", "10.docx", "깨짐.docx"]   # 이름 순
        assert [x["상태"] for x in s["줄"]] == ["완료", "완료", "완료", "실패"]
        assert s["파일있음"]
        한장 = Document(io.BytesIO(urllib.request.urlopen(U + "/download/1").read()))
        assert [p.text for p in 한장.paragraphs if re.match(r"후보\d+ ", p.text)] == [
            "후보1 책임연구원", "후보2 책임연구원", "후보10 책임연구원"]
        요약 = Document(io.BytesIO(urllib.request.urlopen(U + "/download/2").read()))
        assert len(요약.tables[0].rows) == 4
    finally:
        srv.shutdown()


# --- 2쪽으로 넘어간 표 ------------------------------------------------------------------
def _두쪽파일(사이글: str = "", 빈칸줄: bool = False, sdt: bool = False) -> bytes:
    """1쪽 표(항목 1~4) + 2쪽 표(머리줄 다시 + 항목 5~7) + 작성자."""
    doc = Document()
    doc.add_paragraph("홍길동 책임연구원")
    doc.add_paragraph("(한국과학기술원)")
    for 쪽, 범위 in ((1, 항목명[:4]), (2, 항목명[4:])):
        if 쪽 == 2:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if 사이글:
                doc.add_paragraph(사이글)
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "검증 의견"
        for 항 in 범위:
            c = t.add_row().cells
            c[0].text, c[1].text = 항, f"{항} 의견 본문"
        if 빈칸줄 and 쪽 == 2:
            c = t.add_row().cells
            c[0].text, c[1].text = "", "앞 항목에서 이어지는 글"
        if sdt and 쪽 == 2:
            from docx.oxml import OxmlElement
            감싸기, 속 = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
            t._tbl.addprevious(감싸기)
            감싸기.append(속)
            속.append(t._tbl)
    doc.add_paragraph("작성자 김민수 / 이지은 / 박서준")
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


@pytest.mark.parametrize("모양", [{}, {"사이글": "- 2 -"}, {"빈칸줄": True}, {"sdt": True}])
def test_2쪽까지_다_읽는다(모양):
    자료 = D.검증자료_읽기(_두쪽파일(**모양))
    assert [a for a, _b in 자료.항목] == 항목명, 모양
    assert 자료.항목[6][1].startswith("종합 의견 의견 본문")
    if 모양.get("빈칸줄"):
        assert 자료.항목[6][1].endswith("앞 항목에서 이어지는 글")
    assert D.작성자_찾기(자료.끝글) == ["김민수", "이지은", "박서준"]
